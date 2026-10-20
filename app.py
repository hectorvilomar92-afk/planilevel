import os
import requests
import streamlit as st
from duckduckgo_search import DDGS
from pptx import Presentation
from pptx.util import Inches, Pt

# Configuración de la interfaz web
st.set_page_config(
    page_title="Generador de Presentaciones Educativas",
    page_icon="🎓",
    layout="centered",
)

# --- SISTEMA DE LOGIN SENCILLO ---
CREDENCIALES = {"profesor": "profe2026", "estudiante": "clase2026"}

if "autenticado" not in st.session_state:
  st.session_state.autenticado = False

if not st.session_state.autenticado:
  st.title("🔐 Acceso a la Plataforma Educativa")
  st.write(
      "Inicia sesión para generar tus presentaciones y recursos didácticos."
  )

  with st.form("login_form"):
    user = st.text_input("Usuario")
    password = st.text_input("Contraseña", type="password")
    entrar = st.form_submit_button("Ingresar")

    if entrar:
      if user in CREDENCIALES and CREDENCIALES[user] == password:
        st.session_state.autenticado = True
        st.session_state.usuario = user
        st.rerun()
      else:
        st.error("Usuario o contraseña incorrectos.")
  st.stop()

# --- APLICACIÓN PRINCIPAL ---
st.sidebar.success(f"Conectado como: **{st.session_state.usuario}**")
if st.sidebar.button("Cerrar Sesión"):
  st.session_state.autenticado = False
  st.rerun()

st.title("🎓 Generador de Presentaciones con Enlaces de YouTube")
st.write(
    "Crea presentaciones automáticas con investigación web y links de videos de"
    " YouTube listos para exponer."
)

with st.form("form_presentacion"):
  tema = st.text_input(
      "¿Sobre qué tema o asignatura deseas la presentación?",
      "La Célula y sus Partes",
  )
  col1, col2 = st.columns(2)
  with col1:
    num_slides = st.slider(
        "Número de diapositivas", min_value=3, max_value=8, value=4
    )
  with col2:
    incluir_yt = st.checkbox(
        "Incluir enlaces de videos de YouTube", value=True
    )

  generar_btn = st.form_submit_button("🚀 Generar y Descargar Presentación")


# Función para buscar información y videos en la web
@st.cache_data(show_spinner=False)
def buscar_contenido_y_videos(query):
  texto_info = ""
  videos_encontrados = []
  try:
    with DDGS() as ddgs:
      res_texto = [r["body"] for r in ddgs.text(query, max_results=5)]
      if res_texto:
        texto_info = " ".join(res_texto)

res_yt = [
    r
    for r in ddgs.text(
        f"site:youtube.com watch?v {query}", max_results=3
    )
]
      for v in res_yt:
        titulo_v = v.get("title", "Video educativo")
        url_v = v.get("href", "#")
        videos_encontrados.append({"titulo": titulo_v, "url": url_v})
  except Exception:
    pass

  return (
      texto_info
      or f"Contenido académico e investigativo de referencia sobre {query}.",
      videos_encontrados,
  )


if generar_btn:
  if not tema.strip():
    st.warning("Por favor, escribe un tema válido.")
  else:
    barra = st.progress(0)
    estado = st.empty()

    try:
      estado.text(
          "🌐 Investigando contenido y buscando videos en YouTube en tiempo"
          " real..."
      )
      barra.progress(30)
      informacion, lista_videos = buscar_contenido_y_videos(tema)

      estado.text("📝 Estructurando las diapositivas...")
      barra.progress(60)

      prs = Presentation()

      # Generar diapositivas de contenido
      for i in range(num_slides):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = f"Sección {i+1}: {tema}"
        slide.placeholders[1].text = (
            f"Desarrollo teórico y análisis del tema.\n\nInformación clave:"
            f" {informacion[:250]}..."
        )

      # Diapositiva final con los links de YouTube
      if incluir_yt and lista_videos:
        slide_yt = prs.slides.add_slide(prs.slide_layouts[1])
        slide_yt.shapes.title.text = "🎥 Enlaces de Videos en YouTube"
        detalle_yt = (
            "Recursos audiovisuales recomendados para complementar la"
            " clase:\n\n"
        )
        for v in lista_videos:
          detalle_yt += f"• {v['titulo']}\n  🔗 {v['url']}\n\n"
        slide_yt.placeholders[1].text = detalle_yt

      barra.progress(90)
      estado.text("📊 Compilando archivo PowerPoint...")

      # Guardar temporalmente en el entorno de ejecución para descarga inmediata
      archivo_salida = f"presentacion_{os.getpid()}.pptx"
      prs.save(archivo_salida)

      barra.progress(100)
      estado.text("¡Completado con éxito!")

      with open(archivo_salida, "rb") as f:
        st.success("🎉 ¡Tu presentación está lista para descargar!")
        st.download_button(
            label="📥 Descargar Presentación (.pptx)",
            data=f,
            file_name=f"{tema.replace(' ', '_')}.pptx",
            mime=(
                "application/vnd.openxmlformats-officedocument.presentationml.presentation"
            ),
        )

      # Eliminar archivo del servidor inmediatamente después de habilitar la descarga
      if os.path.exists(archivo_salida):
        os.remove(archivo_salida)

    except Exception as err:
      st.error(f"Ocurrió un error: {err}")
